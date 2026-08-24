"""Verify the timesheet's formula logic.

LibreOffice is unavailable in this environment, so instead of a recalc-and-eyeball
we build a small instance of the *same* generator, inject known inputs, evaluate
every formula with the `formulas` engine, and assert the numbers.

Also statically checks the full-size workbook for functions that commonly fail
outside Excel and for malformed references.
"""

import datetime as dt
import os
import re
import subprocess
import sys

import formulas
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SMALL = "/tmp/ts_small.xlsx"
FULL = os.path.join(HERE, "Real_Estate_Professional_Timesheet.xlsx")

_BANNED_NAMES = ["XLOOKUP", "XMATCH", "SORT", "FILTER", "UNIQUE", "SEQUENCE",
                 "TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS"]
# Boundary-anchored so SUMIFS/COUNTIFS do not match bare IFS, and an _xlfn.
# prefixed name is not flagged either.
BANNED = [(n, re.compile(r"(?<![A-Z0-9_.])" + n + r"\s*\(")) for n in _BANNED_NAMES]

failures = []
checks = 0


def check(label, actual, expected, tol=1e-6):
    global checks
    checks += 1
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        ok = abs(float(actual) - float(expected)) < tol
    else:
        ok = str(actual).strip() == str(expected).strip()
    if not ok:
        failures.append(f"{label}: expected {expected!r}, got {actual!r}")
    print(f"  {'PASS' if ok else 'FAIL'}  {label}: {actual!r}")


def coerce(v):
    """formulas returns numpy/schedula wrappers; unwrap to a scalar."""
    for _ in range(4):
        if hasattr(v, "value"):
            v = v.value
            continue
        if hasattr(v, "__len__") and not isinstance(v, str):
            try:
                if len(v) == 0:
                    return ""
                v = v[0]
                continue
            except (TypeError, IndexError):
                break
        break
    if hasattr(v, "item"):
        try:
            v = v.item()
        except (ValueError, AttributeError):
            pass
    return v


# ---------------------------------------------------------------- build small
print("Building small instance of the real generator...")
env = dict(os.environ, TS_LOG_ROWS="20", TS_N_PROPS="4", TS_N_CATS="25", TS_OUT=SMALL)
subprocess.run([sys.executable, os.path.join(HERE, "build_timesheet.py")],
               cwd=HERE, env=env, check=True)

# ---------------------------------------------------------------- inject data
wb = load_workbook(SMALL)
lg = wb["Time Log"]
st = wb["Setup"]

OAK = "123 Oak Street - Rental"
MAPLE = "456 Maple Ave - Rental"
ROWS = [
    (dt.date(2025, 1, 6),  "Self",   OAK,   "Tenant screening & leasing", 100.0),
    (dt.date(2025, 1, 10), "Self",   OAK,   "Reviewing financial statements / operating reports", 50.0),
    (dt.date(2025, 2, 5),  "Self",   MAPLE, "Property showings & client tours", 700.0),
    (dt.date(2025, 3, 5),  "Spouse", OAK,   "Rent collection & tenant communications", 450.0),
]
for i, (d, who, prop, cat, hrs) in enumerate(ROWS):
    r = 5 + i
    lg.cell(row=r, column=1, value=d)
    lg.cell(row=r, column=2, value=who)
    lg.cell(row=r, column=3, value=prop)
    lg.cell(row=r, column=4, value=cat)
    lg.cell(row=r, column=5, value="test entry")
    lg.cell(row=r, column=6, value=hrs)
lg.cell(row=5, column=10, value=None)          # drop the example-row marker
st.cell(row=4, column=8, value=200.0)          # January non-real-estate hours
wb.save(SMALL)

# ------------------------------------------------------------ locate by label
probe = load_workbook(SMALL)
db = probe["Dashboard"]


def row_of(label):
    for r in range(1, db.max_row + 1):
        v = db.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip() == label:
            return r
    raise AssertionError(f"label not found on Dashboard: {label!r}")


r_qual = row_of("Your qualifying real property hours")
r_nonq = row_of("Your non-qualifying hours (investor, education, commuting)")
r_nonre = row_of("Your hours in non-real-estate work (from Setup)")
r_allsvc = row_of("Your total personal services in all trades or businesses")
r_spouse = row_of("Spouse hours logged on properties (material participation only)")
r_total = row_of("Total hours logged (both spouses)")
r_t750 = row_of("More than 750 hours in real property trades or businesses")
r_t50 = row_of("More than 50% of your personal services in real property")
r_overall = row_of("OVERALL - both must be met by you alone")
r_gap = row_of("Hours you still need to clear 750")
r_agg = row_of("AGGREGATED - all rentals as one activity (Sec. 469(c)(7)(A) election)")
r_prop0 = row_of("Property / Entity") + 1        # first property row
r_cat0 = row_of("Activity Category") + 1         # first category row
r_mon0 = row_of("Month") + 1                     # January

# ------------------------------------------------------------------- evaluate
print("\nEvaluating model with `formulas`...")
xl = formulas.ExcelModel().loads(SMALL).finish()
sol = xl.calculate()

norm = {}
for k, v in sol.items():
    m = re.match(r"^'\[.*?\](.+?)'!([A-Z]+\d+)$", k)
    if m:
        norm[(m.group(1).upper(), m.group(2))] = v


def cell(sheet, ref):
    v = norm.get((sheet.upper(), ref))
    if v is None:
        raise AssertionError(f"no evaluated value for {sheet}!{ref}")
    return coerce(v)


print("\nSummary of hours")
check("Your qualifying hours (100 + 700)", cell("Dashboard", f"B{r_qual}"), 800)
check("Your non-qualifying hours (50)", cell("Dashboard", f"B{r_nonq}"), 50)
check("Your non-RE hours (200)", cell("Dashboard", f"B{r_nonre}"), 200)
check("Total personal services (800 + 200)", cell("Dashboard", f"B{r_allsvc}"), 1000)
check("Spouse hours (450, excluded from tests)", cell("Dashboard", f"B{r_spouse}"), 450)
check("Total logged both spouses (1300)", cell("Dashboard", f"B{r_total}"), 1300)

print("\nStatutory tests")
check("750-hour actual", cell("Dashboard", f"B{r_t750}"), 800)
check("750-hour result", cell("Dashboard", f"D{r_t750}"), "MET")
check(">50% actual (800/1000)", cell("Dashboard", f"B{r_t50}"), 0.8)
check(">50% result", cell("Dashboard", f"D{r_t50}"), "MET")
check("Overall", cell("Dashboard", f"D{r_overall}"), "QUALIFIES")
check("Gap to 750 (already past)", cell("Dashboard", f"B{r_gap}"), 0)

print("\nQualifying-activity classification on the Time Log")
check("Row 5 category qualifies", cell("Time Log", "G5"), "Yes")
check("Row 6 investor activity excluded", cell("Time Log", "G6"), "No")
check("Row 5 month = 1", cell("Time Log", "H5"), 1)
check("Row 7 month = 2", cell("Time Log", "H7"), 2)

print("\nMaterial participation by property (spouse hours included)")
prop_rows = {}
for i in range(4):
    name = coerce(cell("Dashboard", f"A{r_prop0 + i}"))
    prop_rows[str(name)] = r_prop0 + i
ro = prop_rows[OAK]
rm = prop_rows[MAPLE]
check("Oak - your hours (100 + 50)", cell("Dashboard", f"B{ro}"), 150)
check("Oak - spouse hours (450)", cell("Dashboard", f"C{ro}"), 450)
check("Oak - combined (600)", cell("Dashboard", f"D{ro}"), 600)
check("Oak - 500-hour test", cell("Dashboard", f"E{ro}"), "Met")
check("Maple - your hours (700)", cell("Dashboard", f"B{rm}"), 700)
check("Maple - spouse hours (0)", cell("Dashboard", f"C{rm}"), 0)
check("Maple - 500-hour test", cell("Dashboard", f"E{rm}"), "Met")
check("Brokerage row - below 500", cell("Dashboard", f"E{prop_rows['Brokerage - General (no property)']}"),
      "Below 500")
check("Blank property row stays blank", cell("Dashboard", f"E{r_prop0 + 3}"), "")

print("\nAggregation election row")
check("Aggregate your hours (850)", cell("Dashboard", f"B{r_agg}"), 850)
check("Aggregate spouse hours (450)", cell("Dashboard", f"C{r_agg}"), 450)
check("Aggregate combined (1300)", cell("Dashboard", f"D{r_agg}"), 1300)
check("Aggregate 500-hour test", cell("Dashboard", f"E{r_agg}"), "Met")

print("\nBy month")
check("Jan your total (100 + 50)", cell("Dashboard", f"B{r_mon0}"), 150)
check("Jan your qualifying (100)", cell("Dashboard", f"C{r_mon0}"), 100)
check("Jan spouse (0)", cell("Dashboard", f"D{r_mon0}"), 0)
check("Feb your qualifying (700)", cell("Dashboard", f"C{r_mon0 + 1}"), 700)
check("Mar spouse (450)", cell("Dashboard", f"D{r_mon0 + 2}"), 450)
check("Cumulative through Mar (800)", cell("Dashboard", f"E{r_mon0 + 2}"), 800)

print("\nBy category")
cat_rows = {}
for i in range(25):
    name = coerce(cell("Dashboard", f"A{r_cat0 + i}"))
    cat_rows[str(name)] = r_cat0 + i
rt = cat_rows["Tenant screening & leasing"]
rv = cat_rows["Reviewing financial statements / operating reports"]
check("Tenant screening hours (100)", cell("Dashboard", f"B{rt}"), 100)
check("Tenant screening % of qualifying (100/800)", cell("Dashboard", f"D{rt}"), 0.125)
check("Investor category hours (50)", cell("Dashboard", f"B{rv}"), 50)
check("Investor category flagged No", cell("Dashboard", f"C{rv}"), "No")
check("Investor category has no % (excluded)", cell("Dashboard", f"D{rv}"), "")

# ------------------------------------------------------- static check on FULL
print("\nStatic check of the full-size workbook")
full = load_workbook(FULL)
sheets = {s.upper() for s in full.sheetnames}
n_formulas = 0
static_problems = []

for ws in full.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or not c.value.startswith("="):
                continue
            n_formulas += 1
            f = c.value.upper()
            for name, pat in BANNED:
                if pat.search(f):
                    static_problems.append(f"{ws.title}!{c.coordinate}: banned function {name}")
            if f.count("(") != f.count(")"):
                static_problems.append(f"{ws.title}!{c.coordinate}: unbalanced parentheses")
            for ref in re.findall(r"'([^']+)'!", c.value):
                if ref.upper() not in sheets:
                    static_problems.append(f"{ws.title}!{c.coordinate}: unknown sheet {ref!r}")
            for ref in re.findall(r"(?<!['\w!])([A-Za-z][A-Za-z0-9_]*)!", c.value):
                if ref.upper() not in sheets:
                    static_problems.append(f"{ws.title}!{c.coordinate}: unknown sheet {ref!r}")
                if " " in ref:
                    static_problems.append(f"{ws.title}!{c.coordinate}: unquoted sheet with space")

checks += 1
if static_problems:
    failures.extend(static_problems[:20])
    print(f"  FAIL  {len(static_problems)} static problems")
else:
    print(f"  PASS  {n_formulas} formulas, no banned functions or bad references")

# --------------------------------------------------------------------- report
print("\n" + "=" * 62)
if failures:
    print(f"FAILED - {len(failures)} of {checks} checks")
    for f in failures:
        print(f"  - {f}")
    sys.exit(1)
print(f"ALL {checks} CHECKS PASSED ({n_formulas} formulas in the full workbook)")
